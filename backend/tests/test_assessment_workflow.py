from io import BytesIO
from uuid import UUID, uuid4

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from PIL import Image

from app.database import SessionLocal
from app.models import ExtractionResult, ProcessingJob, Rubric, RubricItem
from conftest import csrf_headers


def png_bytes() -> bytes:
    image = Image.new("RGB", (120, 80), color=(244, 245, 248))
    stream = BytesIO()
    image.save(stream, format="PNG")
    return stream.getvalue()


def prepare_exam(client: TestClient, *, pupil_name: str = "Student"):
    from conftest import register_and_verify

    register_and_verify(
        client, f"assessment-{uuid4()}@example.test", school_name="Assessment School"
    )
    year = client.post(
        "/api/v1/academics/years",
        json={"label": "2026/2027", "is_current": True},
        headers=csrf_headers(client),
    )
    assert year.status_code == 201, year.text
    klass = client.post(
        "/api/v1/academics/classes",
        json={"name": f"SS 2 {uuid4().hex[:5]}"},
        headers=csrf_headers(client),
    )
    subject = client.post(
        "/api/v1/academics/subjects",
        json={"name": f"Biology {uuid4().hex[:5]}"},
        headers=csrf_headers(client),
    )
    term = client.post(
        "/api/v1/academics/terms",
        json={"academic_year_id": year.json()["id"], "name": "First Term"},
        headers=csrf_headers(client),
    )
    assert klass.status_code == subject.status_code == term.status_code == 201
    student = client.post(
        "/api/v1/students",
        json={
            "admission_number": f"ADM-{uuid4().hex[:8]}",
            "first_name": pupil_name,
            "last_name": "Learner",
            "class_id": klass.json()["id"],
            "academic_year_id": year.json()["id"],
        },
        headers=csrf_headers(client),
    )
    assert student.status_code == 201, student.text
    exam = client.post(
        "/api/v1/exams",
        json={
            "title": "Biology assessment",
            "class_id": klass.json()["id"],
            "subject_id": subject.json()["id"],
            "academic_year_id": year.json()["id"],
            "term_id": term.json()["id"],
            "exam_type": "Theory",
            "max_marks": 10,
        },
        headers=csrf_headers(client),
    )
    assert exam.status_code == 201, exam.text
    return student.json(), exam.json()


def test_secure_scan_teacher_approval_result_publication_and_exports(client):
    # This is a deliberately formula-like input to verify exports remain literal strings.
    student, exam = prepare_exam(client, pupil_name='=HYPERLINK("https://example.invalid")')
    valid_items = [
        {
            "question_number": "Q1",
            "prompt": "Explain the concept",
            "expected_answer": "A teacher-defined expected answer",
            "alternatives": ["An accepted equivalent"],
            "criteria": "Partial credit is awarded for supported facts.",
            "max_marks": 10,
        }
    ]
    invalid_total = client.put(
        f"/api/v1/exams/{exam['id']}/rubric",
        json={"items": [{**valid_items[0], "max_marks": 9}]},
        headers=csrf_headers(client),
    )
    assert invalid_total.status_code == 422
    saved = client.put(
        f"/api/v1/exams/{exam['id']}/rubric",
        json={"items": valid_items},
        headers=csrf_headers(client),
    )
    assert saved.status_code == 200, saved.text

    uploaded = client.post(
        "/api/v1/papers",
        data={"exam_id": exam["id"], "student_id": student["id"]},
        files={"file": ("answer.png", png_bytes(), "image/png")},
        headers=csrf_headers(client),
    )
    assert uploaded.status_code == 201, uploaded.text
    paper = uploaded.json()["paper"]
    page_rows = client.get(f"/api/v1/papers/{paper['id']}/pages").json()
    assert len(page_rows) == 1 and page_rows[0]["source_page_number"] == 1
    page_update = client.patch(
        f"/api/v1/papers/{paper['id']}/pages/{page_rows[0]['id']}",
        json={"rotation": 90, "flagged": True},
        headers=csrf_headers(client),
    )
    assert page_update.status_code == 200
    assert page_update.json()["rotation"] == 90 and page_update.json()["flagged"] is True
    preview = client.get(f"/api/v1/papers/{paper['id']}/pages/{page_rows[0]['id']}/preview")
    assert preview.status_code == 200 and preview.headers["content-type"].startswith("image/png")
    assert client.get(f"/api/v1/papers/{paper['id']}/file").status_code == 200

    with SessionLocal() as db:
        rubric = db.query(Rubric).filter_by(exam_id=UUID(exam["id"])).one()
        rubric_item = db.query(RubricItem).filter_by(rubric_id=rubric.id).one()
        db.add(
            ExtractionResult(
                organization_id=rubric.organization_id,
                paper_id=UUID(paper["id"]),
                answers=[
                    {
                        "question_number": "Q1",
                        "answer_text": "Transcribed answer",
                        "page_number": 1,
                        "confidence": 0.86,
                        "source_region": "center",
                        "warning": "Check handwriting",
                    }
                ],
                provider="test-fixture",
                model="local-fixture",
                warnings=["Review this transcription"],
                status="needs_review",
            )
        )
        db.commit()

    correction = client.put(
        f"/api/v1/papers/{paper['id']}/extraction",
        json={
            "answers": [
                {
                    "question_number": "Q1",
                    "answer_text": "Teacher-corrected answer",
                    "page_number": 1,
                    "confidence": 0.86,
                    "source_region": "center",
                    "warning": None,
                }
            ],
            "warnings": [],
        },
        headers=csrf_headers(client),
    )
    assert correction.status_code == 200 and correction.json()["status"] == "needs_review"
    assert (
        client.post(
            f"/api/v1/papers/{paper['id']}/extraction/confirm", headers=csrf_headers(client)
        ).status_code
        == 200
    )
    old_rubric_item_id = rubric_item.id
    with SessionLocal() as db:
        db.get(Rubric, rubric.id).is_locked = True
        db.commit()
    revised = client.put(
        f"/api/v1/exams/{exam['id']}/rubric",
        json={"items": [{**valid_items[0], "criteria": "Revised teacher criteria."}]},
        headers=csrf_headers(client),
    )
    assert revised.status_code == 200 and revised.json()["version"] == 2
    stale_mark = client.post(
        f"/api/v1/papers/{paper['id']}/grades/confirm",
        json={"rubric_item_id": str(old_rubric_item_id), "score": 7},
        headers=csrf_headers(client),
    )
    assert stale_mark.status_code == 404
    with SessionLocal() as db:
        latest_rubric = db.query(Rubric).filter_by(exam_id=UUID(exam["id"])).order_by(Rubric.version.desc()).first()
        rubric_item = db.query(RubricItem).filter_by(rubric_id=latest_rubric.id).one()
    over_max = client.post(
        f"/api/v1/papers/{paper['id']}/grades/confirm",
        json={"rubric_item_id": str(rubric_item.id), "score": 11},
        headers=csrf_headers(client),
    )
    assert over_max.status_code == 422
    confirmed = client.post(
        f"/api/v1/papers/{paper['id']}/grades/confirm",
        json={"rubric_item_id": str(rubric_item.id), "score": 7.5},
        headers=csrf_headers(client),
    )
    assert confirmed.status_code == 200
    locked_page = client.patch(
        f"/api/v1/papers/{paper['id']}/pages/{page_rows[0]['id']}",
        json={"rotation": 180},
        headers=csrf_headers(client),
    )
    assert locked_page.status_code == 409
    locked_assignment = client.patch(
        f"/api/v1/papers/{paper['id']}/assignment",
        json={"student_id": None},
        headers=csrf_headers(client),
    )
    assert locked_assignment.status_code == 409
    results = client.get(f"/api/v1/results?exam_id={exam['id']}")
    assert results.status_code == 200 and results.json()[0]["total_score"] == 7.5
    assert results.json()[0]["status"] == "finalized"
    assert (
        client.post(
            f"/api/v1/results/{exam['id']}/{student['id']}/publish", headers=csrf_headers(client)
        ).json()["status"]
        == "published"
    )
    no_reason = client.post(
        f"/api/v1/papers/{paper['id']}/grades/confirm",
        json={"rubric_item_id": str(rubric_item.id), "score": 8},
        headers=csrf_headers(client),
    )
    assert no_reason.status_code == 422
    correction = client.post(
        f"/api/v1/papers/{paper['id']}/grades/confirm",
        json={
            "rubric_item_id": str(rubric_item.id),
            "score": 8,
            "correction_reason": "Reviewed the source page again",
        },
        headers=csrf_headers(client),
    )
    assert correction.status_code == 200
    assert client.get(f"/api/v1/results?exam_id={exam['id']}").json()[0]["status"] == "finalized"
    actions = [row["action"] for row in client.get("/api/v1/audit").json()]
    assert "exam_result.published" in actions
    assert "exam_result.unpublished_after_mark_correction" in actions

    csv_response = client.get(f"/api/v1/results/export.csv?exam_id={exam['id']}")
    assert csv_response.status_code == 200 and "'=HYPERLINK" in csv_response.text
    xlsx_response = client.get(f"/api/v1/results/export.xlsx?exam_id={exam['id']}")
    assert xlsx_response.status_code == 200
    workbook = load_workbook(BytesIO(xlsx_response.content), read_only=True)
    sheet = workbook.active
    assert sheet is not None
    assert str(sheet.cell(row=2, column=2).value).startswith("'=HYPERLINK")


def test_foreign_school_cannot_access_a_private_exam_paper(client):
    from conftest import register_and_verify

    student, exam = prepare_exam(client)
    uploaded = client.post(
        "/api/v1/papers",
        data={"exam_id": exam["id"], "student_id": student["id"]},
        files={"file": ("scan.png", png_bytes(), "image/png")},
        headers=csrf_headers(client),
    )
    paper_id = uploaded.json()["paper"]["id"]
    with TestClient(client.app) as other:
        register_and_verify(
            other, f"second-school-{uuid4()}@example.test", school_name="Another School"
        )
        assert other.get(f"/api/v1/papers/{paper_id}/file").status_code == 404
        assert other.get(f"/api/v1/papers/{paper_id}/pages").status_code == 404
        assert other.get(f"/api/v1/exams/{exam['id']}").status_code == 404


def test_grading_jobs_are_server_deduplicated_and_rubric_is_snapshotted(client, monkeypatch):
    from types import SimpleNamespace

    from app import tasks

    student, exam = prepare_exam(client)
    items = [{
        "question_number": "Q1",
        "prompt": "Explain the concept",
        "expected_answer": "A teacher-defined expected answer",
        "alternatives": [],
        "criteria": "Credit supported facts.",
        "max_marks": 10,
    }]
    saved = client.put(
        f"/api/v1/exams/{exam['id']}/rubric",
        json={"items": items},
        headers=csrf_headers(client),
    )
    assert saved.status_code == 200
    uploaded = client.post(
        "/api/v1/papers",
        data={"exam_id": exam["id"], "student_id": student["id"]},
        files={"file": ("answer.png", png_bytes(), "image/png")},
        headers=csrf_headers(client),
    )
    assert uploaded.status_code == 201, uploaded.text
    paper_id = UUID(uploaded.json()["paper"]["id"])
    rubric_id = UUID(saved.json()["id"])
    with SessionLocal() as db:
        rubric = db.get(Rubric, rubric_id)
        db.add(ExtractionResult(
            organization_id=rubric.organization_id,
            paper_id=paper_id,
            answers=[{"question_number": "Q1", "answer_text": "Student answer", "page_number": 1, "confidence": 1.0}],
            provider="test-fixture",
            model="local-fixture",
            warnings=[],
            status="confirmed",
        ))
        db.commit()

    monkeypatch.setattr("app.routers.papers.settings", SimpleNamespace(vision_ready=True))
    dispatched = []
    monkeypatch.setattr(tasks, "grade_paper", SimpleNamespace(delay=lambda job_id: dispatched.append(job_id)))
    request_headers = csrf_headers(client)
    first = client.post(
        f"/api/v1/papers/{paper_id}/grade",
        headers={**request_headers, "Idempotency-Key": "click-one"},
    )
    second = client.post(
        f"/api/v1/papers/{paper_id}/grade",
        headers={**request_headers, "Idempotency-Key": "a-new-client-key"},
    )
    third = client.post(f"/api/v1/papers/{paper_id}/grade", headers=request_headers)
    assert first.status_code == second.status_code == third.status_code == 202
    assert first.json()["id"] == second.json()["id"] == third.json()["id"]
    assert dispatched == [first.json()["id"]]
    with SessionLocal() as db:
        jobs = db.query(ProcessingJob).filter_by(paper_id=paper_id, task_type="grade").all()
        assert len(jobs) == 1
        assert jobs[0].rubric_id == rubric_id
        assert db.get(Rubric, rubric_id).is_locked is True
